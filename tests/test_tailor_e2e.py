"""End-to-end: parse -> match -> tailor -> review/edit -> re-score -> export, with Gemini mocked."""

from __future__ import annotations

import io

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from app.match_schema import RelevanceRatings
from app.pipeline import sections
from app.pipeline.run import ResumePipeline
from app.tailor_schema import TailorDraft
from tests.match_fakes import PYTHON_JD_TEXT
from tests.sample_resume import FakeProvider
from tests.tailor_fakes import DISHONEST_DRAFT, INFOSYS_1, FakeTailorProvider
from tests.downloads import follow
from tests.test_render import needs_chromium


def _match(client: TestClient, pdf: bytes) -> dict:
    created = client.post("/match", data={"jd_text": PYTHON_JD_TEXT}, files={"file": ("sample_resume.pdf", pdf, "application/pdf")})
    match = client.get(f"/match/{created.json()['match_id']}").json()
    assert match["status"] == "done", match["error"]
    return match


def _tailor(client: TestClient, match_id: str, **extra) -> dict:
    created = client.post("/tailor", json={"match_id": match_id, "template": "modern", **extra})
    assert created.status_code == 202, created.text
    res = client.get(f"/tailor/{created.json()['tailor_id']}").json()
    assert res["status"] == "done", res["error"]
    return res


@needs_chromium
def test_parse_match_tailor_export_rescore(settings: Settings, sample_pdf_bytes: bytes) -> None:
    provider = FakeTailorProvider()
    with TestClient(create_app(settings, provider=provider)) as client:
        match = _match(client, sample_pdf_bytes)
        tailored = _tailor(client, match["match_id"])

        # Re-score: both scores stored, per-component comparison available.
        assert tailored["old_score"] == match["score"]
        assert isinstance(tailored["new_score"], int)
        assert tailored["score_change"] == tailored["new_score"] - tailored["old_score"]
        assert set(tailored["old_breakdown"]) == set(tailored["new_breakdown"])
        assert provider.count(TailorDraft) == 1 and provider.count(RelevanceRatings) == 2  # match + re-score
        assert 0.2 in provider.temperatures

        doc = tailored["document"]
        assert doc["basics"]["email"] == "priya.sharma@example.com"
        assert tailored["changes_log"] and tailored["placeholders"] == []
        assert tailored["page_count"] == 1 and tailored["target_pages"] in (1, 2)

        # Export both formats with the Firstname_Lastname_<Company>_<Role> name.
        link = client.post(f"/tailor/{tailored['tailor_id']}/export", json={"format": "pdf", "template": "classic", "page_size": "Letter"})
        assert link.json()["file_name"] == "Priya_Sharma_FinStackTechnologies_PythonDeveloperBackend.pdf"
        assert link.json()["page_count"] == 1
        pdf = follow(client, link)
        assert pdf.content.startswith(b"%PDF")
        docx_link = client.post(f"/tailor/{tailored['tailor_id']}/export", json={"format": "docx"})
        assert docx_link.json()["file_name"].endswith(".docx")
        docx = follow(client, docx_link)
        assert "Designed and built REST APIs" in "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)

        # The exported PDF is ATS-readable by our own parser.
        extraction = ResumePipeline(settings, provider=FakeProvider()).extract_text(pdf.content)
        found = sections.detect_sections(extraction.blocks).order
        assert {"summary", "experience", "projects", "skills", "education"} <= set(found)

        # History.
        listing = client.get("/tailors").json()
        assert listing[0]["tailor_id"] == tailored["tailor_id"] and listing[0]["company"] == "FinStack Technologies"

        # Same match again -> the existing tailored resume (no new Gemini call) unless regenerate.
        again = client.post("/tailor", json={"match_id": match["match_id"]}).json()
        assert again == {"tailor_id": tailored["tailor_id"], "status": "done", "existing": True}
        assert provider.count(TailorDraft) == 1


@needs_chromium
def test_placeholders_block_export_until_resolved(settings: Settings, sample_pdf_bytes: bytes) -> None:
    with TestClient(create_app(settings, provider=FakeTailorProvider(draft=DISHONEST_DRAFT))) as client:
        tailored = _tailor(client, _match(client, sample_pdf_bytes)["match_id"])
        tid = tailored["tailor_id"]
        assert {p["placeholder"] for p in tailored["placeholders"]} == {"[X]", "[if true: Kubernetes]", "[X%]"}
        assert any(w.startswith("reverted_company") for w in tailored["warnings"])

        blocked = client.post(f"/tailor/{tid}/export", json={"format": "pdf"})
        assert blocked.status_code == 409
        assert len(blocked.json()["detail"]["placeholders"]) == 4  # summary [X], bullet [X], [if true:], [X%]
        assert client.post(f"/tailor/{tid}/export", json={"format": "docx"}).status_code == 409

        # The user fills in / rejects / hides each one, then saves.
        doc = tailored["document"]
        doc["summary"] = doc["summary"].replace("[X]", "3")
        infosys = doc["experience"][0]["bullets"][0]
        infosys["text"], infosys["decision"] = infosys["text"].replace(" for [X] enterprise clients", ""), "edited"
        doc["experience"][1]["bullets"][0]["decision"] = "rejected"
        doc["projects"][0]["hidden"] = True
        # Attempts to change facts are ignored: the company stays the original.
        doc["experience"][0]["company"] = "Google"
        doc["basics"]["email"] = "someone@else.com"

        saved = client.put(f"/tailor/{tid}", json={"document": doc}).json()
        assert saved["placeholders"] == []
        assert saved["document"]["experience"][0]["company"] == "Infosys Ltd"
        assert saved["document"]["basics"]["email"] == "priya.sharma@example.com"
        assert saved["document"]["experience"][0]["bullets"][0]["text"] == (
            "Designed and built REST APIs in Python and Django serving 2M requests per day"
        )
        assert saved["new_score"] is not None

        ok = follow(client, client.post(f"/tailor/{tid}/export", json={"format": "pdf"}))
        with pymupdf.open(stream=ok.content, filetype="pdf") as d:
            text = " ".join(page.get_text() for page in d)
        assert "Developed internal dashboards with React and JS" in text  # the rejected rewrite -> original
        assert "ResumeRank" not in text and "Kubernetes" not in text


def test_edits_rescore_without_llm(settings: Settings, sample_pdf_bytes: bytes) -> None:
    provider = FakeTailorProvider()
    with TestClient(create_app(settings, provider=provider)) as client:
        tailored = _tailor(client, _match(client, sample_pdf_bytes)["match_id"])
        calls = len(provider.calls)
        doc = tailored["document"]
        for job in doc["experience"]:  # hide every experience bullet -> lower relevance / keywords
            for bullet in job["bullets"]:
                bullet["hidden"] = True
        doc["hidden_sections"] = ["projects"]
        doc["section_order"] = ["summary", "skills", "experience", "education", "certifications", "achievements", "projects"]
        saved = client.put(f"/tailor/{tailored['tailor_id']}", json={"document": doc, "template": "compact"}).json()
        assert len(provider.calls) == calls  # re-scoring edits never calls Gemini
        assert saved["new_score"] < tailored["new_score"]
        assert saved["template"] == "compact"
        assert saved["document"]["section_order"][:2] == ["summary", "skills"]

        preview = client.post(f"/tailor/{tailored['tailor_id']}/preview", json={"document": doc, "template": "classic"})
        assert preview.status_code == 200 and "<h2>Skills</h2>" in preview.text and INFOSYS_1 not in preview.text


@needs_chromium
def test_autofit_brings_resume_to_target(settings: Settings, sample_pdf_bytes: bytes) -> None:
    with TestClient(create_app(settings, provider=FakeTailorProvider())) as client:
        tailored = _tailor(client, _match(client, sample_pdf_bytes)["match_id"])
        tid = tailored["tailor_id"]
        # Make it far too long (the sample has ~6 years: target 2 pages): 110 extra bullets, stored directly.
        doc = tailored["document"]
        extra = [
            {"text": f"Delivered backend feature {i} with Python, Django and PostgreSQL for internal teams",
             "original_text": f"feature {i}", "ai_text": "x", "relevance": 0}
            for i in range(110)
        ]
        doc["experience"][0]["bullets"] += [{k: v for k, v in b.items() if k != "relevance"} for b in extra]
        client.app.state.tailors.update(tid, user_edits=doc)

        fitted = client.post(f"/tailor/{tid}/autofit", json={}).json()
        steps = fitted["autofit_steps"]
        assert steps[0].startswith("switched to the Compact template")
        assert any(s.startswith("hid ") for s in steps)
        assert fitted["template"] == "compact"
        assert fitted["page_count"] <= fitted["target_pages"]
        hidden = sum(b["hidden"] for b in fitted["document"]["experience"][0]["bullets"])
        assert 0 < hidden < 112  # trimmed, but the job keeps bullets


def test_tailor_api_errors(settings: Settings, sample_pdf_bytes: bytes) -> None:
    with TestClient(create_app(settings, provider=FakeTailorProvider(fail_on=TailorDraft))) as client:
        assert client.post("/tailor", json={"match_id": "nope"}).status_code == 404
        match = _match(client, sample_pdf_bytes)
        assert client.post("/tailor", json={"match_id": match["match_id"], "template": "fancy"}).status_code == 400
        tid = client.post("/tailor", json={"match_id": match["match_id"]}).json()["tailor_id"]
        res = client.get(f"/tailor/{tid}").json()
        assert res["status"] == "failed" and "TailorDraft is down" in res["error"]
        assert client.post(f"/tailor/{tid}/export", json={"format": "pdf"}).status_code == 400
        assert client.get("/tailor/does-not-exist").status_code == 404
        # A failed tailoring is not re-used: the next request starts a new one.
        assert client.post("/tailor", json={"match_id": match["match_id"]}).json()["existing"] is False


@pytest.mark.parametrize("fmt", ["txt"])
def test_export_rejects_unknown_format(settings: Settings, sample_pdf_bytes: bytes, fmt: str) -> None:
    with TestClient(create_app(settings, provider=FakeTailorProvider())) as client:
        tid = _tailor(client, _match(client, sample_pdf_bytes)["match_id"])["tailor_id"]
        assert client.post(f"/tailor/{tid}/export", json={"format": fmt}).status_code == 400


def test_web_ui_has_tailor_tab(settings: Settings) -> None:
    with TestClient(create_app(settings)) as client:
        page = client.get("/app").text
        for text in ("✂ Tailor", "Tailor my resume for this job", "My tailored resumes", "/tailor/", "Auto-fit", "Accept all"):
            assert text in page, text
