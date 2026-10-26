"""JD sources for Compare Jobs: URL fetching (mocked HTTP) and .txt / .pdf / .docx files."""

from __future__ import annotations

import asyncio
import io
import json

import httpx
import pymupdf
import pytest
from docx import Document

from app.pipeline.jd_sources import (
    URL_ERROR,
    JDSourceError,
    check_url,
    extract_jd_from_html,
    fetch_jd_from_url,
    jd_text_from_file,
)
from tests.match_fakes import PYTHON_JD_TEXT

JD_PARAGRAPHS = "".join(f"<p>{line}</p>" for line in PYTHON_JD_TEXT.splitlines() if line.strip())
JOB_PAGE = f"""<html><head><title>Python Developer - FinStack</title></head><body>
<nav><a href="/">Home</a> <a href="/jobs">All jobs</a> <a href="/login">Sign in</a></nav>
<main><article><h1>Python Developer (Backend)</h1>{JD_PARAGRAPHS}</article></main>
<footer>© FinStack. Privacy · Terms · Cookies</footer></body></html>"""

LINKEDIN_AUTHWALL = """<html><body><h1>Join LinkedIn</h1><p>Sign in to see this job.</p>
<form><input type="email"><input type="password"></form></body></html>"""


def fetch(url: str, handler) -> str:
    return asyncio.run(fetch_jd_from_url(url, transport=httpx.MockTransport(handler)))


# ---------------------------------------------------------------------------
# HTML extraction
# ---------------------------------------------------------------------------


def test_extracts_main_content_and_drops_navigation() -> None:
    text = extract_jd_from_html(JOB_PAGE, "https://jobs.example.com/python-dev")
    assert "Python Developer (Backend)" in text
    assert "REST APIs" in text and "Docker" in text
    assert "Privacy" not in text and "All jobs" not in text


def test_falls_back_to_main_block_when_readability_keeps_only_the_blurb(monkeypatch: pytest.MonkeyPatch) -> None:
    """Seen on Greenhouse pages: readability returned the "About us" paragraph and dropped the duties."""
    import readability

    blurb = "<p>" + "Acme is one of India's leading fintech companies, building the future of money. " * 6 + "</p>"

    class BlurbOnly:
        def __init__(self, html: str) -> None: ...
        def short_title(self) -> str: return "Associate Business Analyst"
        def summary(self, html_partial: bool = False) -> str: return blurb

    monkeypatch.setattr(readability, "Document", BlurbOnly)
    html = f"<html><body><main>{blurb}<h2>Key Responsibilities</h2>{JD_PARAGRAPHS}</main></body></html>"
    text = extract_jd_from_html(html, "https://job-boards.greenhouse.io/acme/jobs/1")
    assert "Key Responsibilities" in text and "REST APIs" in text


def test_prefers_json_ld_job_posting() -> None:
    posting = {
        "@context": "https://schema.org",
        "@graph": [
            {"@type": "Organization", "name": "Ignored"},
            {
                "@type": "JobPosting",
                "title": "Associate Product Manager",
                "hiringOrganization": {"@type": "Organization", "name": "Swiggy"},
                "description": "<p>Own the roadmap for <b>search</b>.</p><ul><li>3+ years of experience</li>"
                "<li>SQL and A/B testing skills required</li></ul>",
            },
        ],
    }
    html = f'<html><head><script type="application/ld+json">{json.dumps(posting)}</script></head><body>App shell</body></html>'
    text = extract_jd_from_html(html, "https://careers.example.com/apm")
    assert text.startswith("Associate Product Manager - Swiggy")
    assert "- SQL and A/B testing skills required" in text


@pytest.mark.parametrize(
    "html, url",
    [
        (LINKEDIN_AUTHWALL, "https://www.linkedin.com/jobs/view/123"),
        (LINKEDIN_AUTHWALL, "https://careers.example.com/job/1"),  # password form + almost no text
        ("<html><body><p>Hello world</p></body></html>", "https://example.com"),  # not a JD
    ],
)
def test_login_walls_and_non_jd_pages_are_refused(html: str, url: str) -> None:
    with pytest.raises(JDSourceError, match=URL_ERROR):
        extract_jd_from_html(html, url)


# ---------------------------------------------------------------------------
# Fetching (mocked HTTP)
# ---------------------------------------------------------------------------


def test_fetch_success_follows_redirects() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/j/1":
            return httpx.Response(301, headers={"Location": "https://jobs.example.com/python-dev"})
        return httpx.Response(200, html=JOB_PAGE)

    assert "Django" in fetch("https://jobs.example.com/j/1", handler)


def test_fetch_redirect_to_login_is_refused() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if "authwall" in request.url.path:
            return httpx.Response(200, html=LINKEDIN_AUTHWALL)
        return httpx.Response(302, headers={"Location": "https://www.linkedin.com/authwall?trk=job"})

    with pytest.raises(JDSourceError, match=URL_ERROR):
        fetch("https://www.linkedin.com/jobs/view/42", handler)


@pytest.mark.parametrize(
    "response",
    [
        httpx.Response(404, html="<p>Not found</p>"),
        httpx.Response(403, html="<p>Forbidden</p>"),
        httpx.Response(200, content=b"%PDF-1.7", headers={"Content-Type": "application/pdf"}),
    ],
)
def test_fetch_http_errors_and_non_html(response: httpx.Response) -> None:
    with pytest.raises(JDSourceError, match=URL_ERROR):
        fetch("https://jobs.example.com/x", lambda request: response)


def test_fetch_network_error() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("connection refused", request=request)

    with pytest.raises(JDSourceError, match=URL_ERROR):
        fetch("https://jobs.example.com/x", handler)


@pytest.mark.parametrize("url", ["", "ftp://example.com/job", "http://localhost:8000/job", "http://127.0.0.1/job", "http://10.0.0.5/"])
def test_bad_or_private_urls_are_refused(url: str) -> None:
    with pytest.raises(JDSourceError, match=URL_ERROR):
        check_url(url)


def test_url_without_scheme_gets_https() -> None:
    assert check_url("jobs.example.com/apm", resolve=False) == "https://jobs.example.com/apm"


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------


def test_txt_file_utf8_and_cp1252() -> None:
    assert "Python Developer" in jd_text_from_file(PYTHON_JD_TEXT.encode("utf-8"), "jd.txt")
    legacy = ("Salary ₹ not shown – " + PYTHON_JD_TEXT).encode("cp1252", errors="replace")
    assert "Python Developer" in jd_text_from_file(legacy, "jd.TXT")


def test_docx_file_paragraphs_and_tables() -> None:
    doc = Document()
    for line in PYTHON_JD_TEXT.splitlines()[:12]:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Location", "Pune (Hybrid)"
    buf = io.BytesIO()
    doc.save(buf)
    text = jd_text_from_file(buf.getvalue(), "jd.docx")
    assert "Python Developer" in text and "Location | Pune (Hybrid)" in text


def test_pdf_file() -> None:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_textbox(pymupdf.Rect(40, 40, 560, 800), PYTHON_JD_TEXT[:1500], fontsize=9)
    text = jd_text_from_file(pdf.tobytes(), "jd.pdf")
    assert "Python" in text and "REST" in text


@pytest.mark.parametrize(
    "data, name, message",
    [
        (b"hello", "jd.rtf", "must be .txt, .pdf or .docx"),
        (b"too short", "jd.txt", "almost no text"),
        (b"not really a pdf", "jd.pdf", "could not be read"),
        (b"not a zip", "jd.docx", "could not be read"),
    ],
)
def test_bad_files(data: bytes, name: str, message: str) -> None:
    with pytest.raises(JDSourceError, match=message):
        jd_text_from_file(data, name)
