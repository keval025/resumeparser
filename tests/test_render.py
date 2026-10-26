"""Export tests: HTML / PDF / DOCX for every template, page counts, and an ATS re-parse of the PDF."""

from __future__ import annotations

import io
from dataclasses import replace

import pymupdf
import pytest
from docx import Document

from app.config import Settings
from app.match_schema import JobDescription
from app.pipeline import sections
from app.pipeline.jd_parse import clean_jd
from app.pipeline.render import (
    TEMPLATES,
    RenderError,
    export_filename,
    fmt_range,
    render_docx,
    render_html,
    render_pdf,
)
from app.pipeline.run import ResumePipeline
from app.pipeline.tailor import display_date, validate_draft
from app.tailor_schema import TailorDraft, TailoredBullet, TailoredResume
from tests.match_fakes import PYTHON_JD
from tests.sample_resume import FakeProvider
from tests.tailor_fakes import CLEAN_DRAFT, DISHONEST_DRAFT


def _chromium_available() -> bool:
    try:
        render_pdf(TailoredResume(basics={"name": "x"}), "compact")
        return True
    except RenderError:
        return False


needs_chromium = pytest.mark.skipif(not _chromium_available(), reason="run: python -m playwright install chromium")


@pytest.fixture
def doc(settings: Settings, sample_pdf_bytes: bytes) -> TailoredResume:
    outcome = ResumePipeline(settings, provider=FakeProvider()).parse_bytes(sample_pdf_bytes, "sample_resume.pdf")
    jd = clean_jd(JobDescription.model_validate(PYTHON_JD), settings.normalization_file)
    tailored, _ = validate_draft(TailorDraft.model_validate(CLEAN_DRAFT), outcome.data, jd, raw_text=outcome.raw_text)
    return tailored


EXPECTED_TEXT = [
    "Priya Sharma",
    "Summary", "Experience", "Projects", "Skills", "Education", "Certifications", "Achievements",
    "Designed and built REST APIs in Python and Django serving 2M requests per day",
    "Tata Consultancy Services",
    "ResumeRank",
]


def _pdf_text(pdf: bytes) -> str:
    with pymupdf.open(stream=pdf, filetype="pdf") as d:
        return " ".join(" ".join(page.get_text().split()) for page in d)


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------


def test_html_is_ats_friendly(doc: TailoredResume) -> None:
    for template in TEMPLATES:
        html = render_html(doc, template)
        assert "<table" not in html and "<img" not in html and "<svg" not in html
        assert '<a href="mailto:priya.sharma@example.com">' in html
        assert '<a href="https://www.linkedin.com/in/priya-sharma-dev">' in html
        for heading in ("Summary", "Experience", "Projects", "Skills", "Education"):
            assert f"<h2>{heading}</h2>" in html


def test_preview_highlights_placeholders_and_escapes(doc: TailoredResume) -> None:
    doc.experience[0].bullets[0].text = "Served <b>2M</b> requests for [N clients]"
    html = render_html(doc, "modern", preview=True)
    assert '<mark class="ph">[N clients]</mark>' in html
    assert "&lt;b&gt;2M&lt;/b&gt;" in html  # resume text is escaped
    assert "<mark" not in render_html(doc, "modern")  # never in exports


def test_hidden_and_rejected_content(doc: TailoredResume) -> None:
    first = doc.experience[0].bullets[0]
    first.decision = "rejected"
    doc.projects[0].hidden = True
    doc.hidden_sections = ["achievements"]
    html = render_html(doc, "classic")
    assert first.original_text in html and first.ai_text not in html
    assert "ResumeRank" not in html and "Achievements" not in html


def test_section_order_is_followed(doc: TailoredResume) -> None:
    doc.section_order = ["summary", "projects", "experience", "skills", "education", "certifications", "achievements"]
    html = render_html(doc, "modern")
    assert html.index("<h2>Projects</h2>") < html.index("<h2>Experience</h2>")


def test_year_only_dates_stay_year_only(doc: TailoredResume) -> None:
    assert fmt_range(doc.education[0].start, doc.education[0].end) == "2017 – 2021"  # resume says "2017 - 2021"
    assert fmt_range(doc.experience[1].start, doc.experience[1].end) == "Jan 2021 – Jun 2021"  # resume says "Jan 2021"
    assert display_date("2019-01", "Worked 01/2019 to 2020") == "2019-01"


def test_export_filename(doc: TailoredResume) -> None:
    assert export_filename(doc, "FinStack Technologies", "Python Developer (Backend)", "pdf") == (
        "Priya_Sharma_FinStackTechnologies_PythonDeveloperBackend.pdf"
    )
    doc.basics.name = "KEVAL HIRPARA"
    assert export_filename(doc, None, "AI PM", "docx") == "Keval_Hirpara_AIPM.docx"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


@needs_chromium
@pytest.mark.parametrize("template", TEMPLATES)
@pytest.mark.parametrize("page_size", ["A4", "Letter"])
def test_pdf_for_every_template(doc: TailoredResume, template: str, page_size: str) -> None:
    pdf, pages = render_pdf(doc, template, page_size)
    assert pdf.startswith(b"%PDF") and pages == 1
    text = _pdf_text(pdf).casefold()  # CSS upper-cases the headings
    for expected in EXPECTED_TEXT:
        assert expected.casefold() in text, (template, expected)
    with pymupdf.open(stream=pdf, filetype="pdf") as d:
        width_pt = d[0].rect.width
        uris = {link.get("uri") for link in d[0].get_links()}
    assert abs(width_pt - (595.3 if page_size == "A4" else 612)) < 2  # Chromium A4 = 595.9pt
    assert "https://www.linkedin.com/in/priya-sharma-dev" in uris  # links stay clickable


@needs_chromium
def test_page_count_grows_and_compact_is_denser(doc: TailoredResume) -> None:
    long_doc = doc.model_copy(deep=True)
    job = long_doc.experience[0]
    job.bullets += [
        TailoredBullet(text=f"Delivered backend feature number {i} with Python, Django and PostgreSQL for internal teams",
                       original_text="x", ai_text="x")
        for i in range(45)
    ]
    modern = render_pdf(long_doc, "modern")[1]
    compact = render_pdf(long_doc, "compact")[1]
    assert modern >= 2
    assert compact < modern


@needs_chromium
def test_exported_pdf_is_ats_readable(doc: TailoredResume, settings: Settings) -> None:
    """Our own parser must find every section in each exported PDF."""
    pipeline = ResumePipeline(replace(settings), provider=FakeProvider())
    for template in TEMPLATES:
        pdf, _ = render_pdf(doc, template)
        extraction = pipeline.extract_text(pdf)
        found = sections.detect_sections(extraction.blocks).order
        for name in ("summary", "experience", "projects", "skills", "education", "certifications", "achievements"):
            assert name in found, (template, found)
        assert extraction.text.startswith("Priya Sharma")
        assert not extraction.is_scanned


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_for_every_template(doc: TailoredResume, template: str) -> None:
    data = render_docx(doc, template, "Letter")
    document = Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs]
    text = "\n".join(paragraphs)
    for expected in EXPECTED_TEXT:
        assert expected in text, (template, expected)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Summary", "Experience", "Projects", "Skills", "Education", "Certifications", "Achievements"]
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert "Designed and built REST APIs in Python and Django serving 2M requests per day" in bullets
    assert not document.tables and not document.inline_shapes  # ATS: no tables, no images
    assert round(document.sections[0].page_width.inches, 1) == 8.5
    assert b"linkedin.com/in/priya-sharma-dev" in data or "linkedin.com/in/priya-sharma-dev" in text


def test_docx_contains_clickable_links(doc: TailoredResume) -> None:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(render_docx(doc, "modern"))) as z:
        rels = z.read("word/_rels/document.xml.rels").decode()
        body = z.read("word/document.xml").decode()
    assert "https://www.linkedin.com/in/priya-sharma-dev" in rels and 'TargetMode="External"' in rels
    assert "<w:hyperlink" in body


def test_dishonest_draft_still_renders(settings: Settings, sample_pdf_bytes: bytes) -> None:
    outcome = ResumePipeline(settings, provider=FakeProvider()).parse_bytes(sample_pdf_bytes, "s.pdf")
    jd = clean_jd(JobDescription.model_validate(PYTHON_JD), settings.normalization_file)
    tailored, _ = validate_draft(TailorDraft.model_validate(DISHONEST_DRAFT), outcome.data, jd)
    html = render_html(tailored, "compact")
    for invented in ("Google", "Senior Software Engineer", "SDE II", "migrate the monolith"):
        assert invented not in html, invented
