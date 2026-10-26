"""Job description sources for Compare Jobs: a job-posting URL or an uploaded .txt / .pdf / .docx file.

URL fetching
------------
``fetch_jd_from_url`` downloads the page with httpx and extracts the JD text, trying in order:
  1. a schema.org ``JobPosting`` in JSON-LD (Greenhouse, Lever, Workday and most job boards embed one),
  2. readability-lxml's main-content extraction, cleaned with BeautifulSoup,
  3. the page's <main> / <article> / <body> text.
Anything that cannot give a usable JD (network error, non-HTML, login wall such as LinkedIn,
too little text) raises ``JDSourceError`` with the one message the UI shows:
"Couldn't read this link, please paste the JD text".
"""

from __future__ import annotations

import asyncio
import io
import ipaddress
import json
import logging
import re
import socket
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup

from app.pipeline.jd_parse import MAX_JD_CHARS, MIN_JD_CHARS, normalize_jd_text

logger = logging.getLogger(__name__)

URL_ERROR = "Couldn't read this link, please paste the JD text"
FILE_TYPES = (".txt", ".pdf", ".docx")
MAX_JD_FILE_BYTES = 5 * 1024 * 1024
MAX_PAGE_BYTES = 3 * 1024 * 1024
FETCH_TIMEOUT = 15.0
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/126.0 Safari/537.36"
)

# Sites whose job pages are only readable when logged in.
_LOGIN_WALL_HOSTS = ("linkedin.com", "glassdoor.", "naukri.com/mnjuser", "wellfound.com/login")
_LOGIN_HINTS = re.compile(
    r"\b(sign ?in|log ?in|join now|create (an )?account|authwall|enable javascript|captcha|access denied)\b", re.IGNORECASE
)
# Words almost every real JD contains at least a couple of.
_JD_HINTS = re.compile(
    r"\b(responsibilit|requirement|qualification|experience|skills?|role|you will|we are looking|about the job|apply)\w*",
    re.IGNORECASE,
)


class JDSourceError(ValueError):
    """The JD could not be read from the URL / file. The message is shown to the user."""


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------


def jd_text_from_file(data: bytes, file_name: str) -> str:
    """Plain text of an uploaded JD file (.txt, .pdf, .docx)."""
    name = (file_name or "").lower()
    if not name.endswith(FILE_TYPES):
        raise JDSourceError(f"'{file_name}': JD files must be .txt, .pdf or .docx.")
    if len(data) > MAX_JD_FILE_BYTES:
        raise JDSourceError(f"'{file_name}' is too large (max {MAX_JD_FILE_BYTES // (1024 * 1024)} MB).")
    try:
        if name.endswith(".txt"):
            text = _decode(data)
        elif name.endswith(".pdf"):
            import pymupdf

            with pymupdf.open(stream=data, filetype="pdf") as doc:
                text = "\n".join(page.get_text() for page in doc)
        else:
            from docx import Document

            document = Document(io.BytesIO(data))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(lines)
    except JDSourceError:
        raise
    except Exception as exc:  # corrupt / password-protected files
        raise JDSourceError(f"'{file_name}' could not be read ({type(exc).__name__}).") from exc
    text = normalize_jd_text(text)
    if len(text) < MIN_JD_CHARS:
        raise JDSourceError(f"'{file_name}' has almost no text (is it a scanned image?). Paste the JD text instead.")
    return text[:MAX_JD_CHARS]


def _decode(data: bytes) -> str:
    encoding = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
    try:
        return data.decode(encoding)
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


# ---------------------------------------------------------------------------
# URLs
# ---------------------------------------------------------------------------


def check_url(url: str, resolve: bool = True) -> str:
    """Validate a user-supplied URL: http(s) only, and (``resolve``) never a private / local address."""
    url = (url or "").strip()
    if url and not re.match(r"^[a-z]+://", url, re.IGNORECASE):
        url = "https://" + url
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise JDSourceError(URL_ERROR)
    host = parsed.hostname.lower()
    if host == "localhost" or host.endswith(".local"):
        raise JDSourceError(URL_ERROR)
    if not resolve:
        return url
    try:
        addresses = {info[4][0] for info in socket.getaddrinfo(host, None)}
    except (socket.gaierror, UnicodeError):
        return url  # unresolvable here; the fetch reports it
    for address in addresses:
        ip = ipaddress.ip_address(address.split("%")[0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise JDSourceError(URL_ERROR)
    return url


def _is_login_wall(url: str, html: str, text: str) -> bool:
    host_path = (urlparse(url).netloc + urlparse(url).path).lower()
    if any(h in host_path for h in _LOGIN_WALL_HOSTS) and len(text) < 1500:
        return True
    if "authwall" in host_path or "/login" in host_path or "/signin" in host_path:
        return True
    soup_has_password = 'type="password"' in html.lower() or "type='password'" in html.lower()
    return soup_has_password and len(text) < 1500


def _html_to_text(fragment: str) -> str:
    soup = BeautifulSoup(fragment, "lxml")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "form", "button"]):
        tag.decompose()
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for li in soup.find_all("li"):
        li.insert_before("\n- ")
    blocks = soup.find_all(["p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "section", "ul", "ol", "tr"])
    for block in blocks:
        block.insert_after("\n")
    return normalize_jd_text(soup.get_text())


def _json_ld_job(soup: BeautifulSoup) -> str | None:
    """Title + company + description of a schema.org JobPosting, if the page embeds one."""

    def walk(node):
        if isinstance(node, list):
            for item in node:
                yield from walk(item)
        elif isinstance(node, dict):
            kind = node.get("@type")
            if kind == "JobPosting" or (isinstance(kind, list) and "JobPosting" in kind):
                yield node
            for key in ("@graph", "mainEntity", "itemListElement"):
                if key in node:
                    yield from walk(node[key])

    for script in soup.find_all("script", type="application/ld+json"):
        try:
            data = json.loads(script.string or script.get_text() or "")
        except (json.JSONDecodeError, TypeError):
            continue
        for job in walk(data):
            description = _html_to_text(str(job.get("description") or ""))
            if len(description) < MIN_JD_CHARS:
                continue
            org = job.get("hiringOrganization")
            company = org.get("name") if isinstance(org, dict) else (org if isinstance(org, str) else None)
            head = " - ".join(str(p) for p in (job.get("title"), company) if p)
            return f"{head}\n{description}" if head else description
    return None


def extract_jd_from_html(html: str, url: str = "") -> str:
    """Main JD text from a job page's HTML. Raises ``JDSourceError`` if it is not a readable JD."""
    soup = BeautifulSoup(html, "lxml")
    text = _json_ld_job(soup)
    if text:  # a structured JobPosting is trusted as is
        return text[:MAX_JD_CHARS]
    try:
        from readability import Document

        doc = Document(html)
        title = (doc.short_title() or "").strip()
        body = _html_to_text(doc.summary(html_partial=True))
        text = f"{title}\n{body}" if title and title.lower() not in body.lower()[:300] else body
    except Exception:  # readability chokes on some markup; fall back below
        text = ""
    # Readability sometimes keeps only the "About us" blurb and drops the duties / requirements
    # (e.g. Greenhouse pages). Then the page's main block is the better source.
    main = soup.find("main") or soup.find("article") or soup.body
    fallback = _html_to_text(str(main)) if main else ""
    readability_ok = bool(_JD_HINTS.search(text)) and len(text) >= 0.5 * len(fallback)
    if not readability_ok and len(fallback) > len(text) * 0.8:
        text = fallback
    page_text = normalize_jd_text(soup.get_text(" "))
    if _is_login_wall(url, html, text) or len(text) < MIN_JD_CHARS * 4:
        raise JDSourceError(URL_ERROR)
    if not _JD_HINTS.search(text) or (len(text) < 800 and _LOGIN_HINTS.search(page_text[:3000])):
        raise JDSourceError(URL_ERROR)
    return text[:MAX_JD_CHARS]


async def fetch_jd_from_url(url: str, transport: httpx.AsyncBaseTransport | None = None, timeout: float = FETCH_TIMEOUT) -> str:
    """Download a job page and return its JD text. Raises ``JDSourceError(URL_ERROR)`` on any failure."""
    # Tests pass a mock transport: then the host is not resolved via DNS.
    url = await asyncio.to_thread(check_url, url, transport is None)
    try:
        async with httpx.AsyncClient(
            follow_redirects=True,
            timeout=timeout,
            transport=transport,
            headers={"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml", "Accept-Language": "en"},
        ) as client:
            response = await client.get(url)
            final_url = str(response.url)
            if response.status_code >= 400:
                logger.info("jd url fetch failed", extra={"url": url, "status": response.status_code})
                raise JDSourceError(URL_ERROR)
            content_type = response.headers.get("content-type", "")
            if "html" not in content_type and "xml" not in content_type and not content_type.startswith("text/"):
                raise JDSourceError(URL_ERROR)
            html = response.text[:MAX_PAGE_BYTES]
    except httpx.HTTPError as exc:
        logger.info("jd url fetch failed", extra={"url": url, "error": type(exc).__name__})
        raise JDSourceError(URL_ERROR) from exc
    if "text/plain" in content_type:
        text = normalize_jd_text(html)
        if len(text) < MIN_JD_CHARS:
            raise JDSourceError(URL_ERROR)
        return text[:MAX_JD_CHARS]
    return extract_jd_from_html(html, final_url)
