"""Render a ``TailoredResume`` to HTML (Jinja2), PDF (Playwright / Chromium) and DOCX (python-docx).

Both exporters use the same ``RenderModel``, so the PDF and the DOCX have the same structure.
Every template is ATS-friendly: one column, standard headings, no tables / images / icons,
selectable text and clickable links.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from functools import lru_cache
import pymupdf
from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup, escape

from app.config import BASE_DIR
from app.pipeline.tailor import PLACEHOLDER, bullet_text, summary_text, visible_sections
from app.tailor_schema import TailoredResume

TEMPLATE_DIR = BASE_DIR / "app" / "templates" / "resume"
TEMPLATES = ("classic", "modern", "compact")
PAGE_SIZES = ("A4", "Letter")
HEADINGS = {
    "summary": "Summary",
    "experience": "Experience",
    "projects": "Projects",
    "skills": "Skills",
    "education": "Education",
    "certifications": "Certifications",
    "achievements": "Achievements",
}
_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


class RenderError(RuntimeError):
    """PDF rendering is unavailable (Playwright / Chromium not installed) or failed."""


# ---------------------------------------------------------------------------
# Render model (shared by HTML and DOCX)
# ---------------------------------------------------------------------------


@dataclass
class Entry:
    title: str = ""
    subtitle: str = ""
    dates: str = ""
    meta: str = ""
    link: str | None = None
    bullets: list[str] = field(default_factory=list)


@dataclass
class Section:
    key: str
    heading: str
    text: str = ""
    entries: list[Entry] = field(default_factory=list)
    skills: list[tuple[str, str]] = field(default_factory=list)
    items: list[str] = field(default_factory=list)


@dataclass
class RenderModel:
    name: str
    contact: list[tuple[str, str | None]]  # (label, href)
    sections: list[Section]


def fmt_date(value: str | None) -> str:
    if not value:
        return ""
    if value == "present":
        return "Present"
    m = re.fullmatch(r"(\d{4})-(\d{2})", value)
    return f"{_MONTHS[int(m.group(2)) - 1]} {m.group(1)}" if m else value


def fmt_range(start: str | None, end: str | None) -> str:
    s, e = fmt_date(start), fmt_date(end)
    return f"{s} – {e}" if s and e else s or e


def _href(url: str) -> str:
    return url if re.match(r"^[a-z]+:", url, re.I) else "https://" + url


def _link_label(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url).rstrip("/")


def build_render_model(doc: TailoredResume) -> RenderModel:
    b = doc.basics
    contact: list[tuple[str, str | None]] = []
    if b.email:
        contact.append((b.email, f"mailto:{b.email}"))
    if b.phone:
        contact.append((b.phone, None))
    if b.location:
        contact.append((b.location, None))
    for url in [b.links.linkedin, b.links.github, b.links.portfolio, *b.links.other]:
        if url:
            contact.append((_link_label(url), _href(url)))

    sections: list[Section] = []
    for key in visible_sections(doc):
        section = Section(key=key, heading=HEADINGS[key])
        if key == "summary":
            section.text = summary_text(doc).strip()
            if not section.text:
                continue
        elif key == "experience":
            for job in doc.experience:
                section.entries.append(
                    Entry(
                        title=job.title or "",
                        subtitle=", ".join(p for p in (job.company, job.location) if p),
                        dates=fmt_range(job.start, job.end),
                        bullets=[bullet_text(x) for x in job.bullets if not x.hidden],
                    )
                )
        elif key == "projects":
            for project in doc.projects:
                if project.hidden:
                    continue
                section.entries.append(
                    Entry(
                        title=project.name or "",
                        meta=", ".join(project.tech_stack),
                        dates=fmt_range(project.start, project.end),
                        link=_href(project.link) if project.link else None,
                        bullets=[bullet_text(x) for x in project.bullets if not x.hidden],
                    )
                )
        elif key == "skills":
            s = doc.skills
            section.skills = [
                (label, ", ".join(values))
                for label, values in (("Technical", s.technical), ("Tools", s.tools), ("Soft skills", s.soft), ("Languages", s.languages))
                if values
            ]
        elif key == "education":
            for e in doc.education:
                degree = " in ".join(p for p in (e.degree, e.field) if p)
                section.entries.append(
                    Entry(title=degree or e.institution or "", subtitle=e.institution if degree else "", dates=fmt_range(e.start, e.end), meta=e.grade or "")
                )
        elif key == "certifications":
            section.items = [
                " – ".join(p for p in (c.name, ", ".join(x for x in (c.issuer, fmt_date(c.date)) if x)) if p)
                for c in doc.certifications if c.name
            ]
        elif key == "achievements":
            section.items = [a for a in doc.achievements if a]
        if section.text or section.entries or section.skills or section.items:
            sections.append(section)
    return RenderModel(name=b.name or "", contact=contact, sections=sections)


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------


@lru_cache(maxsize=1)
def _env() -> Environment:
    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR), autoescape=select_autoescape(["html", "j2"]))
    env.filters["hl"] = lambda text, preview=False: _highlight(text) if preview else escape(text)
    return env


def _highlight(text: str) -> Markup:
    """Escape, then wrap placeholders in <mark> (preview only)."""
    out, last = [], 0
    for m in PLACEHOLDER.finditer(text):
        out += [escape(text[last:m.start()]), Markup('<mark class="ph">'), escape(m.group(0)), Markup("</mark>")]
        last = m.end()
    out.append(escape(text[last:]))
    return Markup("").join(out)


def render_html(doc: TailoredResume, template: str = "modern", page_size: str = "A4", preview: bool = False) -> str:
    if template not in TEMPLATES:
        raise ValueError(f"Unknown template {template!r} (choose from {', '.join(TEMPLATES)})")
    if page_size not in PAGE_SIZES:
        raise ValueError(f"Unknown page size {page_size!r} (choose from {', '.join(PAGE_SIZES)})")
    css = (TEMPLATE_DIR / f"{template}.css").read_text(encoding="utf-8")
    return _env().get_template("resume.html.j2").render(
        r=build_render_model(doc), css=Markup(css), template=template, page_size=page_size, preview=preview
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_PDF_MARGINS = {
    "classic": {"top": "14mm", "bottom": "14mm", "left": "16mm", "right": "16mm"},
    "modern": {"top": "13mm", "bottom": "13mm", "left": "15mm", "right": "15mm"},
    "compact": {"top": "9mm", "bottom": "9mm", "left": "11mm", "right": "11mm"},
}


def html_to_pdf(html: str, page_size: str = "A4", template: str = "modern") -> bytes:
    """Render HTML to PDF with headless Chromium (Playwright)."""
    try:
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RenderError("PDF export needs Playwright: pip install playwright && python -m playwright install chromium") from exc
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            try:
                page = browser.new_page()
                page.set_content(html, wait_until="load")
                return page.pdf(format=page_size, margin=_PDF_MARGINS[template], print_background=True)
            finally:
                browser.close()
    except PlaywrightError as exc:
        if "Executable doesn't exist" in str(exc):
            raise RenderError("Chromium for Playwright is not installed. Run: python -m playwright install chromium") from exc
        raise RenderError(f"PDF rendering failed: {exc}") from exc


def pdf_page_count(pdf: bytes) -> int:
    with pymupdf.open(stream=pdf, filetype="pdf") as d:
        return d.page_count


def render_pdf(doc: TailoredResume, template: str = "modern", page_size: str = "A4") -> tuple[bytes, int]:
    """(PDF bytes, page count)."""
    pdf = html_to_pdf(render_html(doc, template, page_size), page_size, template)
    return pdf, pdf_page_count(pdf)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

_DOCX_STYLE = {
    #           font,              body pt, name pt, heading pt, accent,    margin cm, space-after pt
    "classic": ("Georgia",         10.5,    20,      11.5,       "222222",  1.6,       3),
    "modern": ("Calibri",          10.5,    22,      12,         "2F4F9F",  1.5,       3),
    "compact": ("Calibri",         9.5,     17,      10.5,       "333333",  1.1,       1),
}


def _add_hyperlink(paragraph, url: str, text: str, color: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, attr, value in (("w:color", "w:val", color), ("w:u", "w:val", "single")):
        el = OxmlElement(tag)
        el.set(qn(attr), value)
        props.append(el)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _bottom_border(paragraph, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def render_docx(doc: TailoredResume, template: str = "modern", page_size: str = "A4") -> bytes:
    """Build a DOCX with the same structure as the HTML / PDF."""
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm, Inches, Pt, RGBColor

    if template not in TEMPLATES:
        raise ValueError(f"Unknown template {template!r}")
    font, body_pt, name_pt, heading_pt, accent, margin_cm, after_pt = _DOCX_STYLE[template]
    model = build_render_model(doc)
    document = Document()

    section = document.sections[0]
    if page_size == "Letter":
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    else:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(margin_cm))
    text_width = section.page_width - section.left_margin - section.right_margin

    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(body_pt)
    normal.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font)
    normal.paragraph_format.space_after = Pt(after_pt)
    normal.paragraph_format.space_before = Pt(0)
    heading = document.styles["Heading 1"]
    heading.font.name, heading.font.size, heading.font.bold = font, Pt(heading_pt), True
    heading.font.color.rgb = RGBColor.from_string(accent)
    heading.paragraph_format.space_before = Pt(8 if template != "compact" else 5)
    heading.paragraph_format.space_after = Pt(3 if template != "compact" else 2)
    bullet_style = document.styles["List Bullet"]
    bullet_style.font.name, bullet_style.font.size = font, Pt(body_pt)
    bullet_style.paragraph_format.space_after = Pt(1 if template != "compact" else 0)

    title = document.add_paragraph()
    run = title.add_run(model.name)
    run.bold, run.font.size = True, Pt(name_pt)
    if template == "modern":
        run.font.color.rgb = RGBColor.from_string(accent)
    title.paragraph_format.space_after = Pt(1)

    contact = document.add_paragraph()
    for i, (label, href) in enumerate(model.contact):
        if i:
            contact.add_run("  |  ")
        if href:
            _add_hyperlink(contact, href, label, accent)
        else:
            contact.add_run(label)

    def entry_line(entry) -> None:
        p = document.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(entry.title)
        r.bold = True
        if entry.subtitle:
            p.add_run(f" — {entry.subtitle}")
        if entry.dates:
            p.add_run(f"\t{entry.dates}")
        if entry.meta or entry.link:
            m = document.add_paragraph()
            m.paragraph_format.space_after = Pt(0)
            m.paragraph_format.keep_with_next = True
            if entry.meta:
                r = m.add_run(entry.meta)
                r.italic = True
            if entry.link:
                if entry.meta:
                    m.add_run("  |  ")
                _add_hyperlink(m, entry.link, _link_label(entry.link), accent)
        for bullet in entry.bullets:
            document.add_paragraph(bullet, style="List Bullet")

    for s in model.sections:
        h = document.add_paragraph(s.heading, style="Heading 1")
        if template != "compact":
            _bottom_border(h, accent if template == "modern" else "999999")
        if s.text:
            document.add_paragraph(s.text)
        for entry in s.entries:
            entry_line(entry)
        for label, values in s.skills:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.add_run(f"{label}: ").bold = True
            p.add_run(values)
        for item in s.items:
            document.add_paragraph(item, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# File names
# ---------------------------------------------------------------------------


def _camel(text: str | None, limit: int = 40) -> str:
    words = re.findall(r"[A-Za-z0-9]+", text or "")
    return "".join(w[:1].upper() + w[1:] for w in words)[:limit]


def export_filename(doc: TailoredResume, company: str | None, role: str | None, ext: str) -> str:
    """Firstname_Lastname_<Company>_<Role>.<ext>"""
    names = re.findall(r"[A-Za-z]+", doc.basics.name or "")
    person = [n.capitalize() for n in (names[:1] + names[-1:] if len(names) > 1 else names)] or ["Resume"]
    parts = [*person, _camel(company), _camel(role)]
    return "_".join(p for p in parts if p) + f".{ext}"
